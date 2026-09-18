"""Local, checksum-bound CSV/XLSX ingestion and deterministic quality reporting.

No network calls, raw-file writes, model training, or inferred price units.
"""

import argparse
from collections import Counter
import csv
from datetime import date, datetime, time
import hashlib
from importlib.metadata import version
import io
import json
from pathlib import Path
import platform
import re
import statistics
from openpyxl import load_workbook

from jsonschema import Draft202012Validator, FormatChecker

from .parsing import normalize_text, parse_price, parse_mileage, parse_engine_with_unit, parse_model_year


ROOT = Path(__file__).resolve().parents[2]
PIPELINE_VERSION = "0.2a.2"
SOURCE_COLUMNS = (
    "car_name", "brand", "car_model", "model_year", "transmission",
    "body_type", "fuel_type", "engine_capacity", "kilometers_run", "price",
)
NUMERIC_FIELDS = {"model_year", "engine_capacity", "kilometers_run", "price"}


def digest(content):
    return hashlib.sha256(content).hexdigest()


def canonical_header(value):
    return (normalize_text(value) or "").casefold().replace(" ", "_")


def table_rows(headers, values_iter):
    if not headers:
        raise ValueError("Table header is missing")
    canonical = [canonical_header(value) for value in headers]
    if not all(canonical) or len(set(canonical)) != len(canonical):
        raise ValueError("Empty or duplicate normalized table headers")
    missing = set(SOURCE_COLUMNS) - set(canonical)
    if missing:
        raise ValueError("Required published columns absent: " + ", ".join(sorted(missing)))
    rows = []
    for values in values_iter:
        raw = {name: values[i] if i < len(values) else None for i, name in enumerate(headers)}
        if len(values) > len(headers):
            raw["_extra_cells"] = values[len(headers):]
        mapped = {name: raw[original] for name, original in zip(canonical, headers)}
        rows.append({"raw": raw, "mapped": mapped, "width_mismatch": len(values) != len(headers)})
    return rows


def load_csv(path, *, encoding="utf-8-sig"):
    content = path.read_bytes()
    reader = csv.reader(io.StringIO(content.decode(encoding), newline=""), strict=True)
    headers = next(reader, None)
    return content, headers, table_rows(headers, reader)


def load_xlsx(path, *, sheet=None):
    """Read a matching worksheet without modifying cells or trusting formulas."""
    content = path.read_bytes()
    with io.BytesIO(content) as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            matches = []
            for worksheet in workbook:
                if sheet is not None and worksheet.title != sheet:
                    continue
                headers = next(worksheet.iter_rows(values_only=True), ())
                if set(SOURCE_COLUMNS).issubset({canonical_header(value) for value in headers}):
                    matches.append(worksheet)
            if len(matches) != 1:
                raise ValueError(f"Expected one matching worksheet; found {len(matches)}. Use --sheet to select one.")
            worksheet = matches[0]
            cells = worksheet.iter_rows()
            headers = [str(cell.value) if cell.value is not None else "" for cell in next(cells)]
            values = []
            for source_row in cells:
                if any(cell.data_type == "f" for cell in source_row):
                    raise ValueError("Formula cells require review; cached formula results are not source values")
                values.append([cell.value.isoformat() if isinstance(cell.value, (date, datetime, time)) else cell.value for cell in source_row])
            return content, headers, table_rows(headers, values), worksheet.title
        finally:
            workbook.close()


def load_source(path, *, encoding="utf-8-sig", sheet=None):
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise ValueError("--sheet only applies to XLSX files")
        return (*load_csv(path, encoding=encoding), None)
    if path.suffix.lower() == ".xlsx":
        return load_xlsx(path, sheet=sheet)
    raise ValueError("Supported dataset formats are .csv and .xlsx; extract ZIP archives first")


def _category(value):
    normalized = normalize_text(value)
    return normalized.casefold() if normalized else None


def normalize_row(row, row_number, receipt):
    source = row["mapped"]
    config = receipt["normalization"]
    parsed = {
        "price": parse_price(source["price"], bare_unit=config.get("bare_price_unit")),
        "kilometers_run": parse_mileage(source["kilometers_run"]),
        "engine_capacity": parse_engine_with_unit(source["engine_capacity"], bare_unit=config.get("bare_engine_unit")),
        "model_year": parse_model_year(source["model_year"]),
    }
    issues = []
    for field, result in parsed.items():
        if result.issue:
            issues.append({"field": field, "reason": result.issue, "severity": "exclude"})
    texts = {field: _category(source[field]) for field in SOURCE_COLUMNS if field not in NUMERIC_FIELDS}
    title_year = re.search(r"\b((?:19|20)\d{2})\s*$", texts["car_name"] or "")
    if title_year and parsed["model_year"].value is not None and int(title_year[1]) != parsed["model_year"].value:
        # Neither field proves manufacturing or registration year; retain both for review.
        issues.append({"field": "model_year", "reason": "title_year_conflict", "severity": "review"})
    for field in ("brand", "car_model"):
        if texts[field] is None:
            issues.append({"field": field, "reason": "missing", "severity": "exclude"})
    if row["width_mismatch"]:
        issues.append({"field": "row", "reason": "column_count_mismatch", "severity": "exclude"})

    # Transparent review heuristics, not learned thresholds or corrected values.
    thresholds = receipt["review_thresholds"]
    for field, minimum, maximum in (
        ("price", thresholds["price_min_bdt"], thresholds["price_max_bdt"]),
        ("kilometers_run", 0, thresholds["mileage_max_km"]),
        ("engine_capacity", thresholds["engine_min_cc"], thresholds["engine_max_cc"]),
        ("model_year", thresholds["model_year_min"], date.fromisoformat(receipt["publication_date"]).year + 1),
    ):
        value = parsed[field].value
        if value is not None and not minimum <= value <= maximum:
            issues.append({"field": field, "reason": "outside_review_range", "severity": "review"})

    key = receipt["dataset_id"] + ":" + receipt["file_sha256"][:16] + ":row:" + str(row_number)
    provenance = {
        "source_id": receipt["dataset_id"],
        "source_record_ref": key,
        "source_url": receipt["source_url"],
        "snapshot_ref": "sha256:" + receipt["file_sha256"],
        "collected_on": receipt.get("collection_date"),
        "collection_date_status": "unknown" if receipt.get("collection_date") is None else "known",
        "raw_values": row["raw"],
        "normalization_version": PIPELINE_VERSION,
    }

    def record(kind, **fields):
        return {
            "schema_version": "0.1.0", "entity_type": kind,
            "id": key + ":" + kind, "market": "BD", "provenance": provenance, **fields,
        }

    vehicle = record(
        "vehicle", make=texts["brand"], model=texts["car_model"],
        model_year=parsed["model_year"].value, manufacture_year=None, variant=None,
        engine_cc=parsed["engine_capacity"].value, fuel_type=texts["fuel_type"],
        hybrid_status=None, transmission=texts["transmission"], body_type=texts["body_type"],
    )
    listing = record("listing", vehicle_id=vehicle["id"], market_category=None, listed_on=None, location=None)
    observation = record(
        "listing_observation", listing_id=listing["id"], observed_on=None,
        currency="BDT", price_type="asking", asking_price_bdt=parsed["price"].value,
        amount_kind="full_asking_price" if parsed["price"].value is not None else "unknown",
        mileage_km=parsed["kilometers_run"].value,
    )
    return {
        "source_row_number": row_number,
        "records": [vehicle, listing, observation],
        "issues": issues,
        "parse_interpretations": {field: result.interpretation for field, result in parsed.items()},
        "exact_duplicate_of": None,
        "normalized_duplicate_of": None,
        "usable_for_historical_baseline": False,
        "normalized_source_values": {**texts, **{field: result.value for field, result in parsed.items()}},
    }


def summarize_numbers(values):
    values = sorted(value for value in values if value is not None)
    if not values:
        return {"count": 0, "min": None, "max": None, "mean": None, "median": None, "p05": None, "p25": None, "p75": None, "p95": None}

    def percentile(p):
        position = (len(values) - 1) * p
        lower = int(position)
        upper = min(lower + 1, len(values) - 1)
        return round(values[lower] + (values[upper] - values[lower]) * (position - lower), 2)

    return {"count": len(values), "min": values[0], "max": values[-1],
            "mean": round(statistics.mean(values), 2), "median": statistics.median(values),
            **{label: percentile(p) for label, p in (("p05", .05), ("p25", .25), ("p75", .75), ("p95", .95))}}


def quality_report(raw_rows, headers, normalized, receipt):
    counts = {}
    for field in ("brand", "car_model", "transmission", "fuel_type", "body_type"):
        values = (row["normalized_source_values"][field] for row in normalized)
        counts[field] = dict(sorted(Counter(value for value in values if value is not None).items()))
    pair_counts = Counter(
        json.dumps([row["normalized_source_values"]["brand"], row["normalized_source_values"]["car_model"]], ensure_ascii=False)
        for row in normalized if row["normalized_source_values"]["brand"] and row["normalized_source_values"]["car_model"]
    )
    rare_limit = receipt["review_thresholds"]["rare_category_count_below"]
    issues = Counter(issue["field"] + ":" + issue["reason"] for row in normalized for issue in row["issues"])
    skew = []
    for field, groups in counts.items():
        total = sum(groups.values())
        if total >= receipt["review_thresholds"]["imbalance_min_rows"] and groups:
            label, count = max(groups.items(), key=lambda item: item[1])
            if count / total >= receipt["review_thresholds"]["imbalance_share"]:
                skew.append({"field": field, "category": label, "count": count, "nonmissing_total": total, "share": count / total})
    return {
        "status": "computed_from_local_file", "dataset_id": receipt["dataset_id"],
        "file_sha256": receipt["file_sha256"], "pipeline_version": PIPELINE_VERSION,
        "raw_row_count": len(raw_rows),
        "usable_row_count": sum(row["usable_for_historical_baseline"] for row in normalized),
        "usable_definition": "No exclude/review issue; complete core numeric fields and make/model; later exact duplicates excluded. Not current-market or model-readiness approval.",
        "duplicate_count": sum(row["exact_duplicate_of"] is not None for row in normalized),
        "normalized_duplicate_count": sum(row["normalized_duplicate_of"] is not None for row in normalized),
        "duplicate_definition": "Rows beyond the first matching record, not verified duplicate physical vehicles.",
        "invalid_price_count": sum(any(i["field"] == "price" and i["severity"] == "exclude" and i["reason"] != "missing" for i in row["issues"]) for row in normalized),
        "suspicious_price_count": sum(any(i["field"] == "price" and i["severity"] == "review" for i in row["issues"]) for row in normalized),
        "missing_price_count": sum(any(i["field"] == "price" and i["reason"] == "missing" for i in row["issues"]) for row in normalized),
        "missing_values_per_column": {name: sum(normalize_text(row["raw"].get(name)) is None for row in raw_rows) for name in headers},
        "actual_column_names": headers,
        "publisher_reported_row_count": receipt["publisher_reported_row_count"],
        "matches_publisher_row_count": len(raw_rows) == receipt["publisher_reported_row_count"],
        "extra_columns": sorted({canonical_header(name) for name in headers} - set(SOURCE_COLUMNS)),
        "unique_brands": sorted(counts["brand"]), "unique_brand_count": len(counts["brand"]),
        "unique_models": sorted(counts["car_model"]), "unique_model_count": len(counts["car_model"]),
        "brand_model_counts": dict(sorted(pair_counts.items())),
        "categorical_counts_all_rows": counts,
        "top_brands": [{"brand": key, "count": value} for key, value in sorted(counts["brand"].items(), key=lambda item: (-item[1], item[0]))[:10]],
        "rare_brands": {key: value for key, value in counts["brand"].items() if value < rare_limit},
        "rare_brand_model_pairs": {key: value for key, value in sorted(pair_counts.items()) if value < rare_limit},
        "numeric_summary_all_parseable_rows": {field: summarize_numbers(row["normalized_source_values"][field] for row in normalized) for field in NUMERIC_FIELDS},
        "numeric_summary_usable_rows": {field: summarize_numbers(row["normalized_source_values"][field] for row in normalized if row["usable_for_historical_baseline"]) for field in NUMERIC_FIELDS},
        "price_interpretation_counts": dict(sorted(Counter(row["parse_interpretations"]["price"] or "unresolved" for row in normalized).items())),
        "issue_counts": dict(sorted(issues.items())),
        "flagged_rows": [{"source_row_number": row["source_row_number"], "issues": row["issues"]} for row in normalized if row["issues"]],
        "categorical_imbalance_flags": skew, "review_thresholds": receipt["review_thresholds"],
        "collection_date": receipt.get("collection_date"),
        "limitations": ["Historical asking-price observations only", "No inferred listing dates", "No confirmed transaction prices", "Numeric distributions include all parseable rows unless labeled usable"],
    }


def ingest(raw_path, receipt, output_dir, *, encoding="utf-8-sig", sheet=None):
    if receipt.get("market") != "BD":
        raise ValueError("Only Bangladesh source receipts are supported")
    if receipt["eligibility_status"] != "APPROVED_FOR_EXPERIMENTATION":
        raise ValueError("Source is not approved for experimentation")
    if not receipt.get("file_sha256") or not receipt.get("original_filename") or not receipt.get("retrieved_on"):
        raise ValueError("Raw receipt incomplete: record actual filename, SHA-256, and retrieval date first")
    date.fromisoformat(receipt["retrieved_on"])
    if raw_path.name != receipt["original_filename"]:
        raise ValueError("Original filename does not match receipt")
    content, headers, rows, actual_sheet = load_source(raw_path, encoding=encoding, sheet=sheet)
    if digest(content) != receipt["file_sha256"]:
        raise ValueError("Raw checksum mismatch; source bytes must not be changed")
    for name in ("bare_price", "bare_engine"):
        if receipt["normalization"].get(name + "_unit") and not receipt["normalization"].get(name + "_unit_evidence"):
            raise ValueError("Explicit evidence required for " + name + "_unit")
    schema_path = ROOT / "data/schemas/market_record.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    normalized, exact_seen, normalized_seen = [], {}, {}
    for number, row in enumerate(rows, 1):
        result = normalize_row(row, number, receipt)
        exact_key = json.dumps(row["raw"], sort_keys=True, ensure_ascii=False)
        if exact_key in exact_seen:
            result["exact_duplicate_of"] = exact_seen[exact_key]
            result["issues"].append({"field": "row", "reason": "exact_duplicate", "severity": "exclude"})
        else:
            exact_seen[exact_key] = number
        # Never equate failed numeric parses solely because they both became null.
        normalized_key = json.dumps({key: result["normalized_source_values"][key] if result["normalized_source_values"][key] is not None else row["mapped"][key] for key in SOURCE_COLUMNS}, sort_keys=True, ensure_ascii=False)
        if normalized_key in normalized_seen and result["exact_duplicate_of"] is None:
            result["normalized_duplicate_of"] = normalized_seen[normalized_key]
            result["issues"].append({"field": "row", "reason": "possible_normalized_duplicate", "severity": "review"})
        normalized_seen.setdefault(normalized_key, number)
        for record in result["records"]:
            validator.validate(record)
        result["usable_for_historical_baseline"] = not result["issues"]
        normalized.append(result)
    report = quality_report(rows, headers, normalized, receipt)
    run_manifest = {
        "receipt": receipt, "encoding": encoding if actual_sheet is None else None,
        "source_format": raw_path.suffix.lower(), "worksheet": actual_sheet, "file_sha256": digest(content),
        "actual_row_count": len(rows), "actual_column_names": headers,
        "pipeline_version": PIPELINE_VERSION, "schema_sha256": digest(schema_path.read_bytes()),
        "pipeline_files_sha256": {name: digest((Path(__file__).parent / name).read_bytes()) for name in ("ingestion.py", "parsing.py")},
        "jsonschema_version": version("jsonschema"), "python_version": platform.python_version(),
        "openpyxl_version": version("openpyxl") if actual_sheet is not None else None,
    }
    # A fresh directory prevents overwriting raw inputs, receipts, or earlier runs.
    output_dir.mkdir(parents=True, exist_ok=False)
    (output_dir / "normalized.jsonl").write_text("".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in normalized), encoding="utf-8")
    for name, data in (("quality_report.json", report), ("run_manifest.json", run_manifest)):
        (output_dir / name).write_text(json.dumps(data, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", required=True, type=Path)
    parser.add_argument("--receipt", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--encoding", default="utf-8-sig")
    parser.add_argument("--sheet", help="Optional XLSX worksheet name")
    args = parser.parse_args()
    receipt = json.loads(args.receipt.read_text(encoding="utf-8"))
    report = ingest(args.raw, receipt, args.output_dir, encoding=args.encoding, sheet=args.sheet)
    print(json.dumps({key: report[key] for key in ("raw_row_count", "usable_row_count", "duplicate_count", "invalid_price_count")}, indent=2))


if __name__ == "__main__":
    main()
