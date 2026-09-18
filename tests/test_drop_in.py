"""Drop-in ingestion tests use temporary synthetic files, never real market data."""

from contextlib import redirect_stdout
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import hashlib

from openpyxl import Workbook

from ml.valuation.drop_in import discover_dataset, prepare_run
from ml.valuation.ingestion import SOURCE_COLUMNS, load_xlsx
from ml.valuation.train import metrics


class DropInTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="synthetic-drop-in-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.raw = self.root / "raw"
        self.raw.mkdir()
        self.row = {"car_name": "Synthetic example", "brand": "test brand", "car_model": "test model",
                    "model_year": 2018, "transmission": "automatic", "body_type": "sedan", "fuel_type": "petrol",
                    "engine_capacity": "1500 cc", "kilometers_run": 65000, "price": "BDT 2350000"}

    def csv_file(self, name="original download.CSV", row=None):
        path = self.raw / name
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=SOURCE_COLUMNS)
            writer.writeheader()
            writer.writerow(row or self.row)
        return path

    def xlsx_file(self, name="original download.xlsx", row=None, second_match=False):
        path = self.raw / name
        book = Workbook()
        book.active.title = "Notes"
        book.active.append(["Synthetic test workbook, not market data"])
        sheet = book.create_sheet("Vehicles")
        sheet.append(list(SOURCE_COLUMNS))
        sheet.append([(row or self.row)[key] for key in SOURCE_COLUMNS])
        if second_match:
            other = book.create_sheet("Alternative")
            other.append(list(SOURCE_COLUMNS))
            other.append([self.row[key] for key in SOURCE_COLUMNS])
        book.save(path)
        book.close()
        return path

    def prepare(self, **kwargs):
        with redirect_stdout(io.StringIO()):
            return prepare_run(self.raw, processed_dir=self.root / "processed", **kwargs)

    def test_no_file_explains_exact_required_source_and_destination(self):
        with self.assertRaises(ValueError) as caught:
            discover_dataset(self.raw)
        message = str(caught.exception)
        self.assertIn("DATASET FOUND: NO", message)
        self.assertIn("10.17632/fmb4xmp4k5.2", message)
        self.assertIn(str(self.raw), message)
        self.assertIn(".csv or .xlsx", message)
        self.assertFalse((self.root / "processed").exists())

    def test_detection_uses_columns_and_accepts_original_uppercase_name(self):
        (self.raw / "cars.csv").write_text("unrelated,data\n1,2\n", encoding="utf-8")
        path = self.csv_file()
        discovered = discover_dataset(self.raw)
        self.assertEqual(discovered[0], path)
        self.assertEqual(discovered[2], list(SOURCE_COLUMNS))

    def test_different_matching_files_require_explicit_selection(self):
        first = self.csv_file("first.csv")
        self.csv_file("second.csv", {**self.row, "kilometers_run": 70000})
        with self.assertRaisesRegex(ValueError, "Multiple distinct"):
            discover_dataset(self.raw)
        self.assertEqual(discover_dataset(self.raw, raw=first)[0], first)

    def test_csv_is_preserved_and_snapshot_does_not_create_false_ambiguity(self):
        source = self.csv_file()
        before = source.read_bytes()
        run = self.prepare()
        self.assertEqual(source.read_bytes(), before)
        manifest = json.loads((run / "run_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["actual_row_count"], 1)
        self.assertEqual(manifest["receipt"]["actual_row_count"], 1)
        self.assertIsNone(manifest["receipt"]["collection_date"])
        snapshot = self.raw / "snapshots" / manifest["file_sha256"] / source.name
        self.assertEqual(snapshot.read_bytes(), before)
        self.assertEqual(discover_dataset(self.raw)[0], source)
        row = json.loads((run / "normalized.jsonl").read_text(encoding="utf-8"))
        self.assertEqual(row["records"][2]["asking_price_bdt"], 2350000)

    def test_xlsx_selects_matching_sheet_and_preserves_numeric_source_values(self):
        path = self.xlsx_file()
        before = path.read_bytes()
        run = self.prepare()
        manifest = json.loads((run / "run_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["worksheet"], "Vehicles")
        self.assertEqual(manifest["source_format"], ".xlsx")
        row = json.loads((run / "normalized.jsonl").read_text(encoding="utf-8"))
        self.assertEqual(row["records"][0]["provenance"]["raw_values"]["model_year"], 2018)
        self.assertEqual(path.read_bytes(), before)

    def test_multiple_xlsx_sheets_require_selection_and_formulas_are_rejected(self):
        path = self.xlsx_file(second_match=True)
        with self.assertRaisesRegex(ValueError, "found 2"):
            load_xlsx(path)
        self.assertEqual(load_xlsx(path, sheet="Vehicles")[3], "Vehicles")
        formula_path = self.xlsx_file("formula.xlsx", {**self.row, "price": "=1000000+200000"})
        with self.assertRaisesRegex(ValueError, "Formula cells"):
            load_xlsx(formula_path)

    def test_unresolved_units_stop_before_creating_a_training_run(self):
        self.csv_file(row={**self.row, "price": "23.5", "engine_capacity": "1500"})
        with self.assertRaisesRegex(ValueError, "source units need review"):
            self.prepare()
        self.assertFalse((self.root / "processed").exists())
        with self.assertRaisesRegex(ValueError, "unit-evidence"):
            self.prepare(price_unit="lakh_BDT", engine_unit="cc")
        run = self.prepare(price_unit="lakh_BDT", engine_unit="cc", unit_evidence="Synthetic test author's convention; not official data evidence")
        row = json.loads((run / "normalized.jsonl").read_text(encoding="utf-8"))
        self.assertEqual(row["records"][2]["asking_price_bdt"], 2350000)
        self.assertEqual(row["records"][0]["engine_cc"], 1500)

    def test_median_absolute_error_is_calculated(self):
        result = metrics([100, 200, 300], [110, 200, 350])
        self.assertEqual(result["median_absolute_error_bdt"], 10)
        self.assertEqual(result["mae_bdt"], 20)

    def test_reviewed_units_apply_only_to_exact_file_checksum(self):
        source = self.csv_file(row={**self.row, "price": "2350000", "engine_capacity": "1500"})
        from ml.valuation.ingestion import ROOT
        manifest = json.loads((ROOT / "data/manifests/mendeley_fmb4xmp4k5_v2.json").read_text(encoding="utf-8"))
        manifest["verified_unit_profile"] = {
            "file_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "price_unit": "BDT", "engine_unit": "cc", "evidence": "Synthetic test-only unit review",
        }
        manifest_path = self.root / "data/manifests/mendeley_fmb4xmp4k5_v2.json"
        manifest_path.parent.mkdir(parents=True)
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        with patch("ml.valuation.drop_in.ROOT", self.root):
            run = self.prepare()
        normalized = json.loads((run / "normalized.jsonl").read_text(encoding="utf-8"))
        self.assertEqual(normalized["records"][2]["asking_price_bdt"], 2350000)
        self.csv_file(row={**self.row, "price": "2400000", "engine_capacity": "1500"})
        with patch("ml.valuation.drop_in.ROOT", self.root):
            with self.assertRaisesRegex(ValueError, "source units need review"):
                self.prepare(raw=source)


if __name__ == "__main__":
    unittest.main()
